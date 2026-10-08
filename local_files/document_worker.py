"""Bounded document conversion child. Never execute macros, formulas or PDF JavaScript."""
import base64
from datetime import date, datetime, time
import io
import json
from pathlib import Path
import sys
import zipfile

MAX_TEXT = 100000


def archive_check(path):
    with zipfile.ZipFile(path) as archive:
        if len(archive.infolist()) > 10000 or sum(item.file_size for item in archive.infolist()) > 128*1024*1024:
            raise ValueError('archive_limit')


def value(item):
    return item.isoformat() if isinstance(item, (datetime, date, time)) else item


def process(request):
    path = Path(request['path']); action = request['action']; options = request['options']
    if not isinstance(options, dict): raise ValueError()
    suffix = path.suffix.lower(); out = io.BytesIO(); warnings = []
    if suffix in ('.docx', '.xlsx', '.xlsm') and path.exists(): archive_check(path)
    start = options.get('offset', 0); limit = options.get('limit', 50)
    if type(start) is not int or type(limit) is not int or start < 0 or not 1 <= limit <= 200:
        raise ValueError()
    if suffix == '.pdf':
        from pypdf import PdfReader, PdfWriter
        pdf = PdfReader(path)
        if pdf.is_encrypted: raise ValueError('encrypted_pdf')
        if action == 'read':
            pages = []
            for index in range(start, min(start+limit, len(pdf.pages))):
                stream = pdf.pages[index].get_contents()
                if stream and len(stream.get_data()) > 16*1024*1024: raise ValueError('pdf_stream_limit')
                text = pdf.pages[index].extract_text() or ''
                pages.append(dict(page=index+1, text=text[:MAX_TEXT], truncated=len(text)>MAX_TEXT,
                                  ocr_needed=not text.strip()))
                if sum(len(p['text']) for p in pages) >= MAX_TEXT: break
            return dict(ok=True, format='pdf', pages=pages, total=len(pdf.pages),
                        next_offset=start+len(pages), truncated=start+len(pages)<len(pdf.pages),
                        note='Scanned pages require OCR; layout/table semantics are not guaranteed.')
        if action != 'pdf_pages': raise ValueError('unsupported_action')
        indices = options.get('pages')
        if not isinstance(indices, list) or not 1 <= len(indices) <= 1000 or any(type(i) is not int or i < 1 or i > len(pdf.pages) for i in indices): raise ValueError()
        writer = PdfWriter()
        for index in indices: writer.add_page(pdf.pages[index-1])
        writer.write(out)
        warnings.append('Page selection/reorder may invalidate signatures or omit document-level metadata/forms. Original retained.')
    elif suffix == '.docx':
        from docx import Document
        doc = Document(path) if path.exists() else Document()
        if action == 'read':
            paragraphs = doc.paragraphs
            selected = [p.text for p in paragraphs[start:start+limit]]
            text = '\n'.join(selected)
            tables = [[[cell.text[:2000] for cell in row.cells[:30]] for row in table.rows[:50]] for table in doc.tables[:10]] if start == 0 else []
            return dict(ok=True, format='docx', text=text[:MAX_TEXT], tables=tables,
                        total=len(paragraphs), next_offset=start+len(selected),
                        truncated=start+len(selected)<len(paragraphs) or len(text)>MAX_TEXT,
                        note='Tables capped at 10 tables/50 rows/30 columns; headers, footers and drawings not extracted.')
        if action == 'docx_create':
            if path.exists(): raise ValueError('create_requires_missing')
            paragraphs = options.get('paragraphs', [])
            if not isinstance(paragraphs, list) or len(paragraphs)>1000 or not all(isinstance(p,str) for p in paragraphs): raise ValueError()
            for paragraph in paragraphs: doc.add_paragraph(paragraph)
        elif action == 'docx_replace':
            old = options.get('old_text'); new = options.get('new_text')
            if not isinstance(old, str) or not old or not isinstance(new, str): raise ValueError()
            paragraphs = list(doc.paragraphs)
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells: paragraphs.extend(cell.paragraphs)
            if sum(p.text.count(old) for p in paragraphs) != 1:
                raise ValueError('unique_single_run_match_required')
            # Preserve formatting by replacing within a single run; reject ambiguous or spanning matches.
            hits = [(run, run.text.count(old)) for p in paragraphs for run in p.runs if old in run.text]
            if sum(count for _,count in hits) != 1: raise ValueError('unique_single_run_match_required')
            run = hits[0][0]; run.text = run.text.replace(old, new, 1)
        else: raise ValueError('unsupported_action')
        doc.save(out)
        warnings.append('python-docx may not preserve every unsupported Word feature; compare retained original.')
    elif suffix in ('.xlsx', '.xlsm'):
        import re
        from openpyxl import Workbook, load_workbook
        from openpyxl.utils.cell import range_boundaries
        if suffix == '.xlsm' and not path.exists(): raise ValueError('xlsm_template_required')
        book = load_workbook(path, keep_vba=suffix=='.xlsm', data_only=False) if path.exists() else Workbook()
        sheet_name = options.get('sheet', book.sheetnames[0])
        if not isinstance(sheet_name, str) or sheet_name not in book.sheetnames: raise ValueError('unknown_sheet')
        sheet = book[sheet_name]
        if action == 'read':
            region = options.get('range', 'A1:J50')
            c1,r1,c2,r2 = range_boundaries(region)
            if not all(type(i) is int and i>0 for i in (c1,r1,c2,r2)) or c2<c1 or r2<r1 or (r2-r1+1)*(c2-c1+1)>10000: raise ValueError('cell_range_limit')
            rows = [[value(cell.value) for cell in row] for row in sheet.iter_rows(min_row=r1,max_row=r2,min_col=c1,max_col=c2)]
            result = dict(ok=True,format='xlsx',sheets=book.sheetnames,sheet=sheet_name,range=region,rows=rows,
                          note='Formulas returned as strings; no recalculation, macro execution or external-link refresh.')
            book.close(); return result
        if action != 'xlsx_write': raise ValueError('unsupported_action')
        cells = options.get('cells')
        if not isinstance(cells, dict) or not 1 <= len(cells) <= 10000: raise ValueError()
        for coordinate, item in cells.items():
            if not isinstance(coordinate,str) or not re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]{0,6}',coordinate): raise ValueError()
            c1,r1,c2,r2 = range_boundaries(coordinate)
            if c1>16384 or r1>1048576: raise ValueError('cell_coordinate_limit')
            if item is not None and not isinstance(item,(str,int,float,bool)): raise ValueError()
            if isinstance(item,str) and len(item)>32767: raise ValueError('excel_string_limit')
            sheet[coordinate] = item
        book.save(out); book.close()
        # Confirm serialization preserved requested cell values, not only that it produced a ZIP.
        verified=load_workbook(io.BytesIO(out.getvalue()),keep_vba=suffix=='.xlsm',data_only=False)
        try:
            for coordinate,item in cells.items():
                if verified[sheet_name][coordinate].value != item:
                    raise ValueError('spreadsheet_serialization_mismatch')
        finally: verified.close()
        warnings.append('Unsupported Excel features may change; formulas are not recalculated. Original retained.')
    elif action == 'read' and suffix in ('.png','.jpg','.jpeg','.webp','.gif','.tiff','.bmp'):
        from PIL import Image
        with Image.open(path) as picture:
            picture.load()
            original = picture.size
            picture.thumbnail((1600,1600))
            picture.convert('RGB').save(out, format='PNG')
        return dict(ok=True, format='image', width=original[0], height=original[1],
                    image=base64.b64encode(out.getvalue()).decode(), mime_type='image/png',
                    note='Preview resized to at most 1600 pixels; first frame only.')
    else: raise ValueError('unsupported_document_format')
    if len(out.getvalue()) > 8*1024*1024: raise ValueError('document_write_limit')
    return dict(ok=True, binary=base64.b64encode(out.getvalue()).decode(), warnings=warnings,
                verify_options=options if suffix in ('.xlsx','.xlsm') else {})


def main():
    try:
        raw = sys.stdin.buffer.read(12*1024*1024+1)
        if len(raw)>12*1024*1024: raise ValueError('request_limit')
        result = process(json.loads(raw))
        # Bound protocol payload, including document text and image previews.
        encoded = json.dumps(result, ensure_ascii=False)
        if len(encoded.encode()) > 12*1024*1024: raise ValueError('result_limit')
    except Exception as exc:
        code = str(exc) if isinstance(exc,ValueError) and str(exc).isidentifier() else 'document_processing_failed'
        encoded = json.dumps(dict(ok=False,code=code,recovery='Use a supported unencrypted format or narrow the requested pages/cells; no target write attempted.'))
    print(encoded)


if __name__ == '__main__': main()
