import asyncio
import base64
import io
import json
import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from local_files.extended import Engine
from local_files.execution import run_program
from local_files.filesystem import snapshot
from local_files.server import create_server


class ExtendedTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name).resolve()
        self.fs=Engine(self.root/'state');self.path=self.root/'sample.txt'
    def tearDown(self): self.temp.cleanup()

    def test_large_resume_commit_and_restore(self):
        self.path.write_bytes(b'original')
        version=snapshot(self.path)['version'];data=b'a'*(8*1024*1024)+b'ending'
        begin=self.fs.call('write_begin',path=str(self.path),expected_version=version,total_bytes=len(data))
        token=begin['upload_id']
        first=self.fs.call('write_chunk',upload_id=token,offset=0,content=base64.b64encode(data[:-6]).decode())
        self.assertTrue(first['ok'],first);self.assertEqual(self.path.read_bytes(),b'original')
        resumed=Engine(self.root/'state')
        status=resumed.call('write_status',upload_id=token);self.assertEqual(status['next_offset'],len(data)-6)
        duplicate=resumed.call('write_chunk',upload_id=token,offset=0,content=base64.b64encode(data[:-6]).decode())
        self.assertTrue(duplicate['duplicate'])
        resumed.call('write_chunk',upload_id=token,offset=len(data)-6,content='ending',encoding='utf-8')
        result=resumed.call('write_commit',upload_id=token)
        self.assertTrue(result['verified'],result);self.assertEqual(self.path.read_bytes(),data)
        restored=resumed.call('restore',operation_id=token,expected_version=result['after']['version'])
        self.assertTrue(restored['verified']);self.assertEqual(self.path.read_bytes(),b'original')

    def test_upload_conflict_preserves_target(self):
        self.path.write_text('old');version=snapshot(self.path)['version']
        token=self.fs.call('write_begin',path=str(self.path),expected_version=version,total_bytes=3)['upload_id']
        self.fs.call('write_chunk',upload_id=token,offset=0,content='new',encoding='utf-8')
        self.path.write_text('external')
        result=self.fs.call('write_commit',upload_id=token)
        self.assertEqual(result['code'],'version_conflict');self.assertEqual(self.path.read_text(),'external')
        self.assertTrue(self.fs.call('write_abort',upload_id=token)['aborted'])

    def test_upload_tamper_and_incomplete(self):
        token=self.fs.call('write_begin',path=str(self.path),expected_version='missing',total_bytes=4)['upload_id']
        self.assertEqual(self.fs.call('write_commit',upload_id=token)['code'],'upload_incomplete')
        self.fs.call('write_chunk',upload_id=token,offset=0,content='abcd',encoding='utf-8')
        Path(self.fs.load(token)['backup']).write_text('ABCD')
        self.assertEqual(self.fs.call('write_commit',upload_id=token)['code'],'stage_changed')
        self.assertFalse(self.path.exists())

    def test_upload_retry_after_receipt_failure(self):
        token=self.fs.call('write_begin',path=str(self.path),expected_version='missing',total_bytes=3)['upload_id']
        with patch.object(self.fs,'save',side_effect=OSError('receipt')):
            failed=self.fs.call('write_chunk',upload_id=token,offset=0,content='abc',encoding='utf-8')
        self.assertFalse(failed['ok'])
        retried=self.fs.call('write_chunk',upload_id=token,offset=0,content='abc',encoding='utf-8')
        self.assertEqual(retried['next_offset'],3)

    def test_docx_replace_preserves_run_and_restore(self):
        from docx import Document
        path=self.root/'sample.docx';doc=Document();run=doc.add_paragraph().add_run('Hello world');run.bold=True;doc.save(path)
        read=self.fs.call('document',path=str(path));self.assertIn('Hello',read['text'])
        changed=self.fs.call('document',path=str(path),action='docx_replace',expected_version=read['version'],options=dict(old_text='world',new_text='team'))
        self.assertTrue(changed.get('verified'),changed);self.assertTrue(Document(path).paragraphs[0].runs[0].bold)
        restored=self.fs.call('restore',operation_id=changed['operation_id'],expected_version=changed['after']['version'])
        self.assertTrue(restored['verified']);self.assertEqual(Document(path).paragraphs[0].text,'Hello world')

    def test_spreadsheet_cells_formula_and_conflict(self):
        path=self.root/'sample.xlsx'
        created=self.fs.call('document',path=str(path),action='xlsx_write',expected_version='missing',options=dict(cells={'A1':'Name','B1':12,'B2':'=B1*2'}))
        self.assertTrue(created.get('verified'),created)
        read=self.fs.call('document',path=str(path),options=dict(range='A1:B2'))
        self.assertEqual(read['rows'][1][1],'=B1*2')
        path.write_bytes(path.read_bytes()+b'external')
        changed=self.fs.call('document',path=str(path),action='xlsx_write',expected_version=read['version'],options=dict(cells={'B1':13}))
        self.assertEqual(changed['code'],'version_conflict')

    def test_pdf_read_and_page_selection(self):
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
        path=self.root/'sample.pdf';writer=PdfWriter()
        page=writer.add_blank_page(width=300,height=300)
        font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
        page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):font})})
        content=DecodedStreamObject();content.set_data(b'BT /F1 12 Tf 10 100 Td (Hello PDF) Tj ET')
        page[NameObject('/Contents')]=writer._add_object(content)
        writer.add_blank_page(width=300,height=300);writer.write(path)
        read=self.fs.call('document',path=str(path));self.assertIn('Hello PDF',read['pages'][0]['text'])
        changed=self.fs.call('document',path=str(path),action='pdf_pages',expected_version=read['version'],options=dict(pages=[1]))
        self.assertTrue(changed.get('verified'),changed);self.assertEqual(changed['document']['total'],1)

    def test_native_image_content_protocol(self):
        from PIL import Image
        from mcp import Client
        path=self.root/'image.png';Image.new('RGB',(30,20),(255,0,0)).save(path)
        async def run():
            async with Client(create_server(self.root/'state')) as client:
                result=await client.call_tool('local_image_preview',{'path':str(path)})
                self.assertFalse(result.is_error,result)
                self.assertTrue(any(item.type=='image' for item in result.content))
                names={tool.name for tool in (await client.list_tools()).tools}
                self.assertEqual(len(names),26);self.assertNotIn('local_program_run',names)
            async with Client(create_server(self.root/'enabled',execution=True)) as client:
                names={tool.name for tool in (await client.list_tools()).tools}
                self.assertEqual(len(names),27);self.assertIn('local_program_run',names)
        asyncio.run(run())

    def test_execution_reread_environment_and_timeout(self):
        with patch.dict(os.environ,{'OPENAI_API_KEY':'do-not-pass'}):
            result=run_program(self.root/'state',sys.executable,['-c',"import os;from pathlib import Path;Path('result.txt').write_text('done');print(os.environ.get('OPENAI_API_KEY','absent'))"],str(self.root))
        self.assertTrue(result['ok'],result);self.assertEqual(self.root.joinpath('result.txt').read_text(),'done');self.assertIn('absent',result['output'])
        timed=run_program(self.root/'state',sys.executable,['-c','import time;time.sleep(10)'],str(self.root),timeout=1)
        self.assertEqual(timed['termination'],'timeout');self.assertFalse(timed['ok'])

    def test_execution_no_elevation_and_output_cap(self):
        rejected=run_program(self.root/'state','/usr/bin/sudo',['true'],str(self.root))
        self.assertFalse(rejected['ok'])
        result=run_program(self.root/'state',sys.executable,['-c',"print('x'*300000)"],str(self.root))
        self.assertTrue(result['truncated']);self.assertLessEqual(len(result['output']),192*1024)

    def test_list_does_not_hash_large_tree(self):
        for index in range(5): (self.root/str(index)).write_text('data')
        first=self.fs.call('list_directory',path=str(self.root),limit=2)
        self.assertTrue(first['truncated']);self.assertEqual(len(first['entries']),2)

    def test_excel_rejects_silent_string_truncation(self):
        path=self.root/'long.xlsx'
        result=self.fs.call('document',path=str(path),action='xlsx_write',expected_version='missing',options=dict(cells={'A1':'x'*40000}))
        self.assertFalse(result['ok']);self.assertEqual(result['code'],'excel_string_limit');self.assertFalse(path.exists())

    def test_docx_ambiguous_cross_run_match_rejected(self):
        from docx import Document
        path=self.root/'ambiguous.docx';doc=Document();doc.add_paragraph('TARGET')
        p=doc.add_paragraph();p.add_run('TAR');p.add_run('GET');doc.save(path)
        version=snapshot(path)['version']
        result=self.fs.call('document',path=str(path),action='docx_replace',expected_version=version,options=dict(old_text='TARGET',new_text='updated'))
        self.assertFalse(result['ok']);self.assertEqual(snapshot(path)['version'],version)

    def test_document_reread_conflict_is_not_success(self):
        from docx import Document
        path=self.root/'race.docx';doc=Document();doc.add_paragraph('old');doc.save(path)
        original=self.fs.write
        def raced(*args,**kwargs):
            result=original(*args,**kwargs)
            changed=Document();changed.add_paragraph('external');changed.save(path)
            return result
        with patch.object(self.fs,'write',side_effect=raced):
            result=self.fs.call('document',path=str(path),action='docx_replace',expected_version=snapshot(path)['version'],options=dict(old_text='old',new_text='new'))
        self.assertFalse(result['ok']);self.assertEqual(result['code'],'post_write_conflict')


if __name__=='__main__': unittest.main()
